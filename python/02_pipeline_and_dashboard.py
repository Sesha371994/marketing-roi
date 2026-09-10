"""
Marketing Campaign ROI Dashboard — full pipeline in one script.
"""
import sqlite3
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter

conn = sqlite3.connect("marketing.db")
raw = pd.read_csv("raw_data/campaign_data.csv")
raw.to_sql("raw_campaigns", conn, if_exists="replace", index=False)

with open("sql/cleaning.sql") as f:
    conn.executescript(f.read())
conn.commit()

clean = pd.read_sql("SELECT * FROM clean_campaigns", conn)
print(f"Raw rows: {len(raw)} -> Clean rows: {len(clean)}")

def parse_date(s):
    if pd.isna(s):
        return pd.NaT
    for fmt in ["%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"]:
        try:
            return pd.to_datetime(s, format=fmt)
        except (ValueError, TypeError):
            continue
    return pd.NaT

clean["date"] = clean["date_raw"].apply(parse_date)
clean = clean.drop(columns=["date_raw"])
clean["month"] = clean["date"].dt.strftime("%Y-%m")

# Key marketing metrics (use np.nan for safe division/rounding)
clean["impressions"] = clean["impressions"].replace(0, np.nan)
clean["clicks_safe"] = clean["clicks"].replace(0, np.nan)

clean["ctr"] = (clean["clicks"] / clean["impressions"] * 100).round(2)
clean["conversion_rate"] = (clean["conversions"] / clean["clicks_safe"] * 100).round(2)
clean["cac"] = (clean["spend"] / clean["conversions"].replace(0, np.nan)).round(2)  # Cost per Acquisition
clean["roas"] = (clean["revenue"] / clean["spend"]).round(2)  # Return on Ad Spend
clean = clean.drop(columns=["clicks_safe"])

clean.to_csv("exports/campaign_clean.csv", index=False)

total_spend = clean["spend"].sum()
total_revenue = clean["revenue"].sum()
overall_roas = total_revenue / total_spend
avg_ctr = clean["ctr"].mean()
total_conversions = clean["conversions"].sum()

print(f"Total Spend: ₹{total_spend:,.0f} | Total Revenue: ₹{total_revenue:,.0f}")
print(f"Overall ROAS: {overall_roas:.2f}x | Avg CTR: {avg_ctr:.2f}%")

channel_summary = clean.groupby("channel").agg(
    spend=("spend", "sum"), revenue=("revenue", "sum")
)
channel_summary["roas"] = (channel_summary["revenue"] / channel_summary["spend"]).round(2)
print("\nROAS by Channel:\n", channel_summary.sort_values("roas", ascending=False))

# ---------------- EXCEL DASHBOARD ----------------
wb = Workbook()
FONT = "Arial"
HEADER_FILL = PatternFill(start_color="7B2D8E", end_color="7B2D8E", fill_type="solid")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name=FONT, bold=True, size=14, color="7B2D8E")
LABEL_FONT = Font(name=FONT, bold=True, size=11)
NORMAL_FONT = Font(name=FONT, size=10)

ws_data = wb.active
ws_data.title = "Raw_Data"
cols = ["campaign_id", "channel", "campaign_name", "spend", "clicks",
        "conversions", "revenue", "ctr", "roas", "month"]
for c, col in enumerate(cols, 1):
    cell = ws_data.cell(1, c, col.replace("_", " ").upper())
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
for r, row in enumerate(clean[cols].itertuples(index=False), 2):
    for c, val in enumerate(row, 1):
        cell = ws_data.cell(r, c, val)
        cell.font = NORMAL_FONT
        if cols[c-1] in ("spend", "revenue"):
            cell.number_format = '"₹"#,##0.00'
for c, col in enumerate(cols, 1):
    ws_data.column_dimensions[get_column_letter(c)].width = max(14, len(col) + 4)
n_rows = len(clean) + 1
ws_data.freeze_panes = "A2"

ws = wb.create_sheet("Dashboard")
ws["B2"] = "Marketing Campaign ROI Dashboard"
ws["B2"].font = TITLE_FONT

ch_col = get_column_letter(cols.index("channel") + 1)
sp_col = get_column_letter(cols.index("spend") + 1)
rv_col = get_column_letter(cols.index("revenue") + 1)
mo_col = get_column_letter(cols.index("month") + 1)

r_ch = f"Raw_Data!${ch_col}$2:${ch_col}${n_rows}"
r_sp = f"Raw_Data!${sp_col}$2:${sp_col}${n_rows}"
r_rv = f"Raw_Data!${rv_col}$2:${rv_col}${n_rows}"
r_mo = f"Raw_Data!${mo_col}$2:${mo_col}${n_rows}"

ws["B6"] = "Total Spend"
ws["D6"] = "Total Revenue"
ws["F6"] = "Overall ROAS"
ws["H6"] = "Avg CTR"
for c in ["B6", "D6", "F6", "H6"]:
    ws[c].font = LABEL_FONT

ws["B7"] = f'=SUM({r_sp})'
ws["B7"].number_format = '"₹"#,##0'
ws["D7"] = f'=SUM({r_rv})'
ws["D7"].number_format = '"₹"#,##0'
ws["F7"] = '=D7/B7'
ws["F7"].number_format = '0.00"x"'
ws["H7"] = f'=AVERAGE(Raw_Data!${get_column_letter(cols.index("ctr")+1)}$2:${get_column_letter(cols.index("ctr")+1)}${n_rows})'
ws["H7"].number_format = '0.00"%"'
for c in ["B7", "D7", "F7", "H7"]:
    ws[c].font = Font(name=FONT, size=16, bold=True, color="7B2D8E")

ws["B10"] = "ROAS by Channel"
ws["B10"].font = LABEL_FONT
ws["B11"] = "Channel"
ws["C11"] = "Spend"
ws["D11"] = "Revenue"
ws["E11"] = "ROAS"
for c in ["B11", "C11", "D11", "E11"]:
    ws[c].font = HEADER_FONT
    ws[c].fill = HEADER_FILL

channels = sorted(clean["channel"].unique().tolist())
for i, ch in enumerate(channels, 12):
    ws[f"B{i}"] = ch
    ws[f"C{i}"] = f'=SUMIF({r_ch},B{i},{r_sp})'
    ws[f"D{i}"] = f'=SUMIF({r_ch},B{i},{r_rv})'
    ws[f"E{i}"] = f'=IFERROR(D{i}/C{i},0)'
    ws[f"C{i}"].number_format = '"₹"#,##0'
    ws[f"D{i}"].number_format = '"₹"#,##0'
    ws[f"E{i}"].number_format = '0.00"x"'
last_ch_row = 11 + len(channels)

ws["G10"] = "Monthly Spend vs Revenue"
ws["G10"].font = LABEL_FONT
ws["G11"] = "Month"
ws["H11"] = "Spend"
ws["I11"] = "Revenue"
for c in ["G11", "H11", "I11"]:
    ws[c].font = HEADER_FONT
    ws[c].fill = HEADER_FILL

months = sorted(clean["month"].dropna().unique().tolist())
for i, m in enumerate(months, 12):
    ws[f"G{i}"] = m
    ws[f"H{i}"] = f'=SUMIF({r_mo},G{i},{r_sp})'
    ws[f"I{i}"] = f'=SUMIF({r_mo},G{i},{r_rv})'
    ws[f"H{i}"].number_format = '"₹"#,##0'
    ws[f"I{i}"].number_format = '"₹"#,##0'
last_mo_row = 11 + len(months)

bar = BarChart()
bar.title = "ROAS by Channel"
data_ref = Reference(ws, min_col=5, min_row=11, max_row=last_ch_row)
cats_ref = Reference(ws, min_col=2, min_row=12, max_row=last_ch_row)
bar.add_data(data_ref, titles_from_data=True)
bar.set_categories(cats_ref)
bar.width = 14
bar.height = 8
ws.add_chart(bar, "B18")

line = LineChart()
line.title = "Monthly Spend vs Revenue"
data_ref2 = Reference(ws, min_col=8, max_col=9, min_row=11, max_row=last_mo_row)
cats_ref2 = Reference(ws, min_col=7, min_row=12, max_row=last_mo_row)
line.add_data(data_ref2, titles_from_data=True)
line.set_categories(cats_ref2)
line.width = 14
line.height = 8
ws.add_chart(line, "G18")

for col, w in [("B",16),("C",14),("D",14),("E",12),("F",12),("G",12),("H",14),("I",14)]:
    ws.column_dimensions[col].width = w

wb.save("exports/Marketing_ROI_Dashboard.xlsx")
print("\nSaved: exports/Marketing_ROI_Dashboard.xlsx")
conn.close()
